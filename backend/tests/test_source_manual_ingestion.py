"""Source hardening Gate 3 — manual file ingestion (CSV / XLSX → Parquet assets).

Locks: named upload limits (bytes, rows, columns, cell size, sheets, zip-bomb),
deterministic unique headers, corrupt / .xls refusal with a safe message,
CSV encodings, staged → bound asset lifecycle (ownership, cleanup on failure
and on source delete), traversal-safe storage keys, config never holding rows,
and idempotent lossless migration of legacy inline-row sources.
"""
from __future__ import annotations

import asyncio
import io
import uuid
import zipfile

import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.dialects.postgresql import JSONB, UUID
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from starlette.datastructures import Headers, UploadFile

from app.core.config import settings
from app.core.database import Base
from app.models.manual_source_asset import ManualSourceAsset
from app.models.models import DataSource, DataSourceType, SyncJob
from app.schemas import DataSourceCreate, DataSourceUpdate
from app.services.datasource_crud_service import DataSourceCRUDService
from app.services.manual_assets import service as assets_service
from app.services.manual_assets.parsing import ManualUploadError, parse_csv, parse_xlsx, unique_headers
from app.services.manual_assets.storage import (
    InvalidStorageKey, LocalFileStorage, set_storage_override, validate_storage_key,
)
from app.services.manual_table_connector import create_manual_table_connector


@compiles(JSONB, "sqlite")
def _jsonb_on_sqlite(_type, _compiler, **_kw):
    return "JSON"


@compiles(UUID, "sqlite")
def _uuid_on_sqlite(_type, _compiler, **_kw):
    return "CHAR(36)"


OWNER = uuid.uuid4()
OTHER = uuid.uuid4()


@pytest.fixture()
def storage(tmp_path):
    st = LocalFileStorage(tmp_path / "manual_assets")
    set_storage_override(st)
    yield st
    set_storage_override(None)


@pytest.fixture()
def db(monkeypatch, storage):
    from cryptography.fernet import Fernet

    monkeypatch.setattr(settings, "DATASOURCE_ENCRYPTION_KEY", Fernet.generate_key().decode(), raising=False)
    from app.models.dataset import DatasetTable

    tables = [DataSource.__table__, SyncJob.__table__, ManualSourceAsset.__table__, DatasetTable.__table__]
    for table in tables:
        for col in table.columns:
            d = col.server_default
            if d is not None and "::" in str(getattr(d, "arg", "")):
                monkeypatch.setattr(col, "server_default", None)
    engine = create_engine("sqlite://", future=True, poolclass=StaticPool,
                           connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine, tables=tables)
    Session = sessionmaker(bind=engine, future=True)
    import app.core.database as database

    monkeypatch.setattr(database, "SessionLocal", Session)
    s = Session()
    yield s
    s.close()


def _files(storage):
    return sorted(p.name for p in storage.root.rglob("*.parquet")) if storage.root.exists() else []


def _xlsx(sheets: dict) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class _Req:
    def __init__(self, n=None):
        self.headers = {"content-length": str(n)} if n is not None else {}


class _User:
    def __init__(self, uid=OWNER):
        self.id = uid


def _upload(db, content: bytes, filename: str, user=None):
    from app.api.datasources import parse_manual_file

    f = UploadFile(file=io.BytesIO(content), filename=filename, headers=Headers({}))
    return asyncio.run(parse_manual_file(request=_Req(), file=f, db=db, current_user=user or _User()))


def _err(exc_info):
    d = exc_info.value.detail
    return exc_info.value.status_code, d.get("code"), d.get("message")


# ── limits ───────────────────────────────────────────────────────────────────

def test_file_over_byte_limit_is_refused_413(db, monkeypatch, storage):
    monkeypatch.setattr(settings, "MANUAL_UPLOAD_MAX_BYTES", 1000)
    with pytest.raises(HTTPException) as ei:
        _upload(db, b"a,b\n" + b"1,2\n" * 500, "big.csv")
    assert _err(ei)[:2] == (413, "file_too_large")
    assert _files(storage) == []


def test_too_many_rows(monkeypatch):
    monkeypatch.setattr(settings, "MANUAL_UPLOAD_MAX_ROWS", 3)
    with pytest.raises(ManualUploadError) as ei:
        parse_csv(b"a\n1\n2\n3\n4\n", "x.csv")
    assert ei.value.code == "too_many_rows"
    assert parse_csv(b"a\n1\n2\n3\n", "x.csv")["x"]["rows"][-1] == {"a": "3"}


def test_too_many_columns(monkeypatch):
    monkeypatch.setattr(settings, "MANUAL_UPLOAD_MAX_COLUMNS", 3)
    with pytest.raises(ManualUploadError) as ei:
        parse_csv(b"a,b,c,d\n1,2,3,4\n", "x.csv")
    assert ei.value.code == "too_many_columns"


def test_oversized_cell(monkeypatch):
    monkeypatch.setattr(settings, "MANUAL_UPLOAD_MAX_CELL_CHARS", 10)
    with pytest.raises(ManualUploadError) as ei:
        parse_csv(b"a\n" + b"x" * 50 + b"\n", "x.csv")
    assert ei.value.code == "cell_too_large"
    with pytest.raises(ManualUploadError) as ei:
        parse_xlsx(io.BytesIO(_xlsx({"S": [["a"], ["y" * 50]]})))
    assert ei.value.code == "cell_too_large"


def test_too_many_total_cells(monkeypatch):
    monkeypatch.setattr(settings, "MANUAL_UPLOAD_MAX_TOTAL_CELLS", 10)
    with pytest.raises(ManualUploadError) as ei:
        parse_csv(b"a,b\n" + b"1,2\n" * 10, "x.csv")
    assert ei.value.code == "too_many_cells"


def test_too_many_sheets(monkeypatch):
    monkeypatch.setattr(settings, "MANUAL_UPLOAD_MAX_SHEETS", 2)
    data = _xlsx({"A": [["x"], [1]], "B": [["x"], [1]], "C": [["x"], [1]]})
    with pytest.raises(ManualUploadError) as ei:
        parse_xlsx(io.BytesIO(data))
    assert ei.value.code == "too_many_sheets"


def test_zip_bomb_ratio_refused_before_openpyxl(monkeypatch):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/worksheets/sheet1.xml", b"\0" * (20 * 1024 * 1024))
    import openpyxl

    def _boom(*a, **k):
        raise AssertionError("openpyxl must not open a zip bomb")

    monkeypatch.setattr(openpyxl, "load_workbook", _boom)
    with pytest.raises(ManualUploadError) as ei:
        parse_xlsx(io.BytesIO(buf.getvalue()))
    assert ei.value.code == "archive_ratio"


def test_zip_uncompressed_total_refused(monkeypatch):
    monkeypatch.setattr(settings, "MANUAL_UPLOAD_MAX_UNCOMPRESSED_BYTES", 1000)
    with pytest.raises(ManualUploadError) as ei:
        parse_xlsx(io.BytesIO(_xlsx({"S": [["a"]] + [[i] for i in range(200)]})))
    assert ei.value.code == "archive_too_large"


# ── headers / values ─────────────────────────────────────────────────────────

def test_duplicate_and_blank_headers_are_unique_and_deterministic():
    assert unique_headers(["id", "", "id", "column_2", None, "ID"], 6) == [
        "id", "column_2_2", "id_2", "column_2", "column_5", "ID_3",
    ]
    out = parse_csv(b"name,name,\nA,B,C\n", "d.csv")["d"]
    assert [c["name"] for c in out["columns"]] == ["name", "name_2", "column_3"]
    assert out["rows"] == [{"name": "A", "name_2": "B", "column_3": "C"}]


def test_value_under_blank_header_is_kept():
    out = parse_csv(b"a\n1,extra\n", "w.csv")["w"]
    assert out["rows"] == [{"a": "1", "column_2": "extra"}]


def test_xlsx_values_and_types():
    import datetime as dt

    out = parse_xlsx(io.BytesIO(_xlsx({"S": [["n", "d", "s"], [1, dt.datetime(2024, 1, 2), "x"], [2.5, dt.datetime(2024, 1, 3, 4, 5, 6), "y"]]})))["S"]
    assert out["rows"][0] == {"n": 1, "d": "2024-01-02", "s": "x"}
    assert out["rows"][1]["d"] == "2024-01-03 04:05:06"
    assert {c["name"]: c["type"] for c in out["columns"]} == {"n": "number", "d": "date", "s": "string"}


# ── corrupt / unsupported ────────────────────────────────────────────────────

def test_corrupt_xlsx_is_clean_422_without_internals(db, storage):
    with pytest.raises(HTTPException) as ei:
        _upload(db, b"PK\x03\x04 definitely not a workbook", "bad.xlsx")
    status, code, msg = _err(ei)
    assert status == 422 and code == "invalid_file"
    assert "Traceback" not in msg and "zipfile" not in msg.lower() and "BadZip" not in msg
    assert _files(storage) == []


def test_xls_is_rejected_everywhere(db):
    with pytest.raises(HTTPException) as ei:
        _upload(db, b"\xd0\xcf\x11\xe0legacy", "old.xls")
    assert _err(ei)[:2] == (400, "unsupported_type")
    from app.services.dashboard_html_import_service import parse_uploaded_source_sheets

    with pytest.raises(ValueError):
        parse_uploaded_source_sheets(file_bytes=b"\xd0\xcf\x11\xe0", filename="old.xls")


def test_frontend_no_longer_advertises_xls():
    import pathlib

    root = pathlib.Path(__file__).resolve().parents[2] / "frontend" / "src" / "components"
    for rel in ("datasources/DataSourceForm.tsx", "dashboards/DashboardHtmlImportModal.tsx"):
        text = (root / rel).read_text(encoding="utf-8")
        assert ".xls," not in text and ".xls\"" not in text, rel  # accept= lists
        assert "'xlsx', 'xls'" not in text, rel  # client-side allow-list


# ── CSV encodings ────────────────────────────────────────────────────────────

@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16", "cp1252"])
def test_csv_encodings(encoding):
    text = "name;city\nJosé;Zürich\n" if encoding == "cp1252" else "tên,thành phố\nNguyễn,Hà Nội\n"
    out = parse_csv(text.encode(encoding), "e.csv")["e"]
    first = out["rows"][0]
    if encoding == "cp1252":
        assert first == {"name": "José", "city": "Zürich"}
    else:
        assert first == {"tên": "Nguyễn", "thành phố": "Hà Nội"}


# ── staged upload → bound source ─────────────────────────────────────────────

def _create_source(db, sheets_cfg, owner=OWNER, name="m"):
    return DataSourceCRUDService.create(
        db, DataSourceCreate(name=name, type="manual", config={"sheets": sheets_cfg}), owner_id=owner,
    )


def test_upload_returns_bounded_preview_and_config_holds_no_rows(db, monkeypatch, storage):
    monkeypatch.setattr(settings, "MANUAL_UPLOAD_PREVIEW_ROWS", 5)
    body = "id,v\n" + "".join(f"{i},{i * 2}\n" for i in range(100))
    res = _upload(db, body.encode(), "data.csv")
    sheet = res["sheets"]["data"]
    assert sheet["row_count"] == 100 and len(sheet["preview_rows"]) == 5 and sheet["preview_truncated"]
    assert "rows" not in sheet
    ds = _create_source(db, {"data": {"asset_id": sheet["asset_id"], "columns": [], "row_count": 1}})
    cfg = ds.config
    assert set(cfg["sheets"]["data"]) == {"asset_id", "columns", "row_count"}
    assert cfg["sheets"]["data"]["row_count"] == 100  # server value, not the client's
    assert "rows" not in str(cfg)
    data = create_manual_table_connector(cfg).get_sheet_data("data")
    assert len(data["rows"]) == 100 and data["rows"][7] == {"id": "7", "v": "14"}
    asset = db.get(ManualSourceAsset, uuid.UUID(sheet["asset_id"]))
    assert asset.datasource_id == ds.id and asset.expires_at is None


def test_foreign_or_bound_asset_cannot_be_reused(db, storage):
    res = _upload(db, b"a\n1\n", "x.csv", user=_User(OTHER))
    aid = res["sheets"]["x"]["asset_id"]
    with pytest.raises(ValueError):
        _create_source(db, {"x": {"asset_id": aid}}, owner=OWNER, name="steal")
    mine = _upload(db, b"a\n1\n", "y.csv")["sheets"]["y"]["asset_id"]
    _create_source(db, {"y": {"asset_id": mine}}, name="first")
    with pytest.raises(ValueError):
        _create_source(db, {"y": {"asset_id": mine}}, name="second")


def test_expired_staged_asset_is_refused_and_gc_removes_it(db, storage):
    import datetime as dt

    aid = _upload(db, b"a\n1\n", "x.csv")["sheets"]["x"]["asset_id"]
    asset = db.get(ManualSourceAsset, uuid.UUID(aid))
    asset.expires_at = dt.datetime.now(dt.timezone.utc) - dt.timedelta(hours=1)
    db.commit()
    with pytest.raises(ValueError):
        _create_source(db, {"x": {"asset_id": aid}})
    assert len(_files(storage)) == 1
    assert assets_service.gc_expired_staged(db) == 1
    db.commit()
    assert _files(storage) == [] and db.query(ManualSourceAsset).count() == 0


def test_failed_bind_removes_files_written_in_that_transaction(db, storage):
    # One inline sheet (written as a new asset) + one invalid reference -> rollback.
    with pytest.raises(ValueError):
        _create_source(db, {"a": {"columns": [{"name": "x"}], "rows": [{"x": 1}]},
                            "b": {"asset_id": str(uuid.uuid4())}})
    assert _files(storage) == []
    assert db.query(ManualSourceAsset).count() == 0
    assert db.query(DataSource).count() == 0


def test_stage_failure_cleans_up(db, storage, monkeypatch):
    def _fail(*a, **k):
        raise RuntimeError("disk full /secret/path")

    real = assets_service.stage_sheets

    def _stage_then_fail(*a, **k):
        real(*a, **k)
        _fail()

    import app.services.manual_assets.service as svc

    monkeypatch.setattr(svc, "stage_sheets", _stage_then_fail)
    with pytest.raises(HTTPException) as ei:
        _upload(db, b"a\n1\n", "x.csv")
    status, code, msg = _err(ei)
    assert status == 500 and "/secret/path" not in msg
    assert _files(storage) == []


def test_source_delete_removes_assets_and_files(db, storage):
    aid = _upload(db, b"a\n1\n", "x.csv")["sheets"]["x"]["asset_id"]
    ds = _create_source(db, {"x": {"asset_id": aid}})
    assert len(_files(storage)) == 1
    assert DataSourceCRUDService.delete(db, ds.id)
    assert _files(storage) == []
    assert db.query(ManualSourceAsset).count() == 0


def test_replacing_a_sheet_on_update_deletes_the_old_asset(db, storage):
    a1 = _upload(db, b"a\n1\n", "x.csv")["sheets"]["x"]["asset_id"]
    ds = _create_source(db, {"x": {"asset_id": a1}})
    a2 = _upload(db, b"a\n2\n", "x.csv")["sheets"]["x"]["asset_id"]
    DataSourceCRUDService.update(db, ds.id, DataSourceUpdate(config={"sheets": {"x": {"asset_id": a2}}}), actor_id=OWNER)
    assert len(_files(storage)) == 1
    assert create_manual_table_connector(ds.config).get_sheet_data("x")["rows"] == [{"a": "2"}]


# ── storage keys ─────────────────────────────────────────────────────────────

@pytest.mark.parametrize("key", ["../../etc/passwd", "..\\x.parquet", "/abs.parquet", "ab/../cd.parquet",
                                 "0" * 31 + ".parquet", "G" * 32 + ".parquet", ""])
def test_storage_keys_are_traversal_safe(storage, key):
    with pytest.raises(InvalidStorageKey):
        validate_storage_key(key)
    with pytest.raises(InvalidStorageKey):
        storage.path_for(key)


# ── legacy migration ─────────────────────────────────────────────────────────

LEGACY_ROWS = [
    {"id": 1, "name": "a", "amount": 1.5, "flag": True, "blank": "", "mixed": 1},
    {"id": 2, "name": "", "amount": "", "flag": False, "blank": "", "mixed": "two"},
    {"id": 3, "name": None, "amount": 3, "flag": True, "blank": "", "mixed": None},
]
LEGACY_COLS = [{"name": k, "type": "string"} for k in LEGACY_ROWS[0]]


def _legacy_source(db):
    ds = DataSource(name="legacy", type=DataSourceType.MANUAL, owner_id=OWNER,
                    config={"sheets": {"Sheet1": {"columns": LEGACY_COLS, "rows": LEGACY_ROWS}}})
    db.add(ds)
    db.commit()
    return ds


def test_legacy_source_migrates_idempotently_without_data_loss(db, storage):
    ds = _legacy_source(db)
    before = create_manual_table_connector(dict(ds.config)).get_sheet_data("Sheet1")
    assert before["rows"] == LEGACY_ROWS  # legacy read works before migration

    assert assets_service.migrate_legacy_manual_source(db, ds) is True
    db.refresh(ds)
    assert "rows" not in str(ds.config)
    after = create_manual_table_connector(ds.config).get_sheet_data("Sheet1")
    assert after["rows"] == before["rows"]
    assert [c["name"] for c in after["columns"]] == [c["name"] for c in LEGACY_COLS]

    files = _files(storage)
    assert assets_service.migrate_legacy_manual_source(db, ds) is False  # idempotent
    assert _files(storage) == files
    assert assets_service.migrate_all_legacy(db) == {"migrated": 0, "skipped": 1, "failed": 0}


def test_flat_legacy_config_migrates(db, storage):
    ds = DataSource(name="flat", type=DataSourceType.MANUAL, owner_id=OWNER,
                    config={"columns": [{"name": "a"}], "rows": [{"a": "x"}]})
    db.add(ds)
    db.commit()
    assert assets_service.migrate_all_legacy(db)["migrated"] == 1
    db.refresh(ds)
    assert create_manual_table_connector(ds.config).get_sheet_data("manual_data")["rows"] == [{"a": "x"}]


def test_legacy_migration_failure_keeps_inline_config_and_no_files(db, storage, monkeypatch):
    ds = _legacy_source(db)
    original = dict(ds.config)

    def _fail(*a, **k):
        raise RuntimeError("boom")

    monkeypatch.setattr(assets_service, "encode_parquet", _fail)
    with pytest.raises(RuntimeError):
        assets_service.migrate_legacy_manual_source(db, ds)
    db.refresh(ds)
    assert ds.config == original
    assert _files(storage) == []
    assert create_manual_table_connector(ds.config).get_sheet_data("Sheet1")["rows"] == LEGACY_ROWS


def test_internal_inline_create_is_stored_as_assets(db, storage):
    """The HTML-import wizard creates sources with inline sheets; they must land as assets."""
    ds = _create_source(db, {"S": {"columns": [{"name": "a", "type": "number"}], "rows": [{"a": 1}, {"a": 2}]}})
    assert "rows" not in str(ds.config)
    assert create_manual_table_connector(ds.config).get_sheet_data("S")["rows"] == [{"a": 1}, {"a": 2}]


def test_query_list_tables_and_columns_read_from_assets(db, storage):
    """Every reader (list tables, columns, SQL via DuckDB) goes through the asset path."""
    from app.services.datasource_service import DataSourceConnectionService as C

    aid = _upload(db, b"id,v\n1,10\n2,20\n3,30\n", "sales.csv")["sheets"]["sales"]["asset_id"]
    cfg = _create_source(db, {"sales": {"asset_id": aid}}).config
    assert [t["name"] for t in C.list_tables("manual", cfg)] == ["sales"]
    assert [c["name"] for c in C._sheets_list_columns(cfg, "sales")] == ["id", "v"]
    cols, rows, _ms = C.execute_query("manual", cfg, "SELECT SUM(v) AS total FROM sales")
    assert float(rows[0]["total"]) == 60.0
