"""report_pptx_service — the editable PowerPoint export (no database).

The browser sends what the reader sees; this module only lays it out. The
contract added here: a CHART block without a real picture never ships — a deck
with an empty card where a chart belongs read as a finished report (a real
customer deck had one). KPI, table and text blocks are unaffected.
"""
from __future__ import annotations

import base64
import io

import pytest
from PIL import Image
from pptx import Presentation

from app.services.report_pptx_service import (
    PptxIncompleteError,
    ReportPptxRequest,
    build_report_pptx,
)


def _png(w: int = 120, h: int = 60, colour=(59, 130, 246)) -> str:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), colour).save(buf, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def _deck(tiles: list[dict], *, width: int = 1800, name: str = "Tổng quan") -> Presentation:
    req = ReportPptxRequest(title="Báo cáo điều hành", footer="Báo cáo", pages=[{"name": name, "width": width, "tiles": tiles}])
    return Presentation(io.BytesIO(build_report_pptx(req)))


def _texts(prs) -> list[str]:
    return [s.text_frame.text for slide in prs.slides for s in slide.shapes if s.has_text_frame]


def test_a_valid_chart_image_becomes_a_picture_with_an_editable_title():
    prs = _deck([{"kind": "chart", "x": 0, "y": 0, "w": 600, "h": 300,
                  "title": "SỐ NGƯỜI DÙNG TRUY CẬP TRUNG BÌNH TRÊN NGÀY", "image": _png()}])
    shapes = list(prs.slides[0].shapes)
    assert sum(1 for s in shapes if s.shape_type == 13) == 1
    assert "SỐ NGƯỜI DÙNG TRUY CẬP TRUNG BÌNH TRÊN NGÀY" in _texts(prs)
    assert "(biểu đồ không chụp được)" not in _texts(prs)


@pytest.mark.parametrize("image", [
    None,
    "",
    "not-a-data-url",
    "data:image/png;base64,",
    "data:image/png;base64," + base64.b64encode(b"definitely not a png").decode(),
])
def test_a_chart_without_a_real_picture_refuses_the_deck_and_names_it(image):
    with pytest.raises(PptxIncompleteError) as err:
        _deck([
            {"kind": "kpi", "x": 0, "y": 0, "w": 290, "h": 140, "title": "Doanh thu", "value": "8.0B"},
            {"kind": "chart", "x": 0, "y": 150, "w": 600, "h": 300, "title": "Doanh thu theo vùng", "image": _png()},
            {"kind": "chart", "x": 600, "y": 150, "w": 600, "h": 300, "title": "Xu hướng", "image": image},
        ])
    assert err.value.charts == ["Xu hướng"]


def test_every_chart_without_a_picture_is_named_across_pages():
    req = ReportPptxRequest(title="R", pages=[
        {"name": "A", "width": 1200, "tiles": [{"kind": "chart", "x": 0, "y": 0, "w": 400, "h": 200, "title": "Một", "image": None}]},
        {"name": "B", "width": 1200, "tiles": [{"kind": "chart", "x": 0, "y": 0, "w": 400, "h": 200, "image": None}]},
    ])
    with pytest.raises(PptxIncompleteError) as err:
        build_report_pptx(req)
    assert err.value.charts == ["Một", "B · chart"]


def test_kpi_table_and_text_never_need_a_picture():
    prs = _deck([
        {"kind": "kpi", "x": 0, "y": 0, "w": 290, "h": 140, "title": "Tổng doanh số", "value": "610"},
        {"kind": "text", "x": 300, "y": 0, "w": 600, "h": 140, "text": "Ghi chú: số liệu đã lọc theo Miền Bắc."},
        {"kind": "table", "x": 0, "y": 150, "w": 1190, "h": 300, "title": "Chi tiết giao dịch",
         "columns": ["id", "region", "amount"], "rows": [["1", "North", "100"], ["3", "North", "200"]], "total_rows": 2},
    ])
    texts = _texts(prs)
    for expected in ("Tổng doanh số", "610", "Ghi chú: số liệu đã lọc theo Miền Bắc.", "Chi tiết giao dịch"):
        assert expected in texts, expected
    table = next(s.table for s in prs.slides[0].shapes if s.has_table)
    assert [table.cell(0, c).text for c in range(3)] == ["id", "region", "amount"]
    assert table.cell(2, 1).text == "North"


def test_a_long_table_shows_what_fits_and_states_how_many_rows_are_left_out():
    rows = [[str(i), "North", str(i * 10)] for i in range(25)]
    prs = _deck([{"kind": "table", "x": 0, "y": 0, "w": 1190, "h": 260, "title": "Bảng dài",
                  "columns": ["id", "region", "amount"], "rows": rows, "total_rows": 40}], width=1200)
    table = next(s.table for s in prs.slides[0].shapes if s.has_table)
    shown = len(table.rows) - 1
    assert 1 <= shown < 25
    note = next(t for t in _texts(prs) if t.startswith("… còn"))
    assert note.startswith(f"… còn {40 - shown} dòng không hiển thị")
    for slide in prs.slides:
        for s in slide.shapes:
            assert s.top + s.height <= prs.slide_height, "a block runs off the slide"
